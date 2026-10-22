"""K118 第 4 章：產生 Claude Code 導入商業論證試算表與一頁摘要。

流程：
1. 用 openpyxl 建 business-case.xlsx（所有計算都是 Excel 公式）
2. 用 LibreOffice 無頭模式重算公式（openpyxl 只寫公式、不會算值）
3. 從重算後的檔案讀出結果，產生 BUSINESS-CASE.md（不手抄數字）
4. 用 Python 獨立算一次，和試算表結果交叉比對
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
XLSX_PATH = BASE_DIR / "business-case.xlsx"
MD_PATH = BASE_DIR / "BUSINESS-CASE.md"

BOOK = "書"
ASSUMED = "假設值"

# 假設工作表：(列, 名稱, 值, 單位, 來源)。列號固定，公式靠它引用。
INPUTS = [
    (4, "團隊開發者人數", 10, "人", ASSUMED),
    (5, "每人年度全額成本（薪資＋福利＋稅＋設備）", 180000, "USD/年", ASSUMED),
    (6, "每人每年工作時數", 2000, "小時", ASSUMED),
    (7, "每年工作週數", 50, "週", ASSUMED),
    (8, "高影響任務佔工時比例", 0.40, "%", ASSUMED),
    (9, "高影響任務效率權重", 1.5, "倍", ASSUMED),
    (10, "中影響任務佔工時比例", 0.35, "%", ASSUMED),
    (11, "中影響任務效率權重", 1.0, "倍", ASSUMED),
    (12, "生產力下降期間", 2, "週", BOOK),
    (13, "每人每月訂閱成本", 100, "USD/人/月", ASSUMED),
    (14, "分析期間", 12, "月", ASSUMED),
]

# 情境：(名稱, 效率提升, 培訓時數, 前兩週生產力下降)
# 效率 15/20/30%、培訓 2–4 小時、下降 10–15% 都是書上的範圍；
# 保守情境取最差端、樂觀取最好端、基準取中間值。
SCENARIOS = [
    ("保守", 0.15, 4, 0.15),
    ("基準", 0.20, 3, 0.125),
    ("樂觀", 0.30, 2, 0.10),
]
SCENARIO_COLS = ["B", "C", "D"]

# 計算工作表的列：(列, 名稱, 公式樣板, 數字格式)。{c} 會換成情境欄位。
A = "假設!$C$"
CALC_ROWS = [
    (5, "每小時全額成本", f"={A}5/{A}6", "#,##0.00"),
    (6, "加權係數（高／中影響任務）", f"={A}8*{A}9+{A}10*{A}11", "0.000"),
    (7, "有效效率提升", "={c}2*{c}6", "0.0%"),
    (8, "每月效益（全團隊）", f"={A}4*{A}5/12*{{c}}7", "#,##0"),
    (9, "每月訂閱成本（全團隊）", f"={A}4*{A}13", "#,##0"),
    (10, "每月淨效益", "={c}8-{c}9", "#,##0"),
    (11, "一次性培訓成本", f"={A}4*{{c}}3*{{c}}5", "#,##0"),
    (12, "一次性生產力下降成本", f"={A}4*{A}5/{A}7*{A}12*{{c}}4", "#,##0"),
    (13, "一次性導入成本合計", "={c}11+{c}12", "#,##0"),
    (14, "分析期總效益", f"={{c}}8*{A}14", "#,##0"),
    (15, "分析期總成本", f"={{c}}13+{{c}}9*{A}14", "#,##0"),
    (16, "分析期淨效益", "={c}14-{c}15", "#,##0"),
    (17, "ROI", "=IF({c}15=0,0,{c}16/{c}15)", "0.0%"),
    (18, "回本月數", '=IF({c}10<=0,"不回本",{c}13/{c}10)', "0.00"),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")


def style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def build_inputs_sheet(wb: Workbook) -> None:
    """假設工作表：所有輸入值與來源標註。"""
    ws = wb.active
    ws.title = "假設"
    ws["A1"] = "Claude Code 導入商業論證：輸入假設（黃底可改）"
    ws["A1"].font = Font(bold=True, size=14)
    for col, title in enumerate(["項目", "說明", "值", "單位", "來源"], start=1):
        ws.cell(row=3, column=col, value=title)
    style_header(ws, 3, 5)
    for row, name, value, unit, source in INPUTS:
        ws.cell(row=row, column=1, value=name)
        cell = ws.cell(row=row, column=3, value=value)
        cell.fill = INPUT_FILL
        if unit == "%":
            cell.number_format = "0%"
        elif isinstance(value, int) and value >= 1000:
            cell.number_format = "#,##0"
        ws.cell(row=row, column=4, value=unit)
        ws.cell(row=row, column=5, value=source)
    ws["A16"] = "註：標「假設值」的數字書上沒有，是本練習自訂，導入前請換成公司實際數字。"
    for col, width in zip("ABCDE", [40, 4, 14, 12, 10]):
        ws.column_dimensions[col].width = width


def build_calc_sheet(wb: Workbook) -> None:
    """計算工作表：三種情境，全部用公式。"""
    ws = wb.create_sheet("計算")
    ws["A1"] = "情境"
    for col, (name, *_rest) in zip(SCENARIO_COLS, SCENARIOS):
        ws[f"{col}1"] = name
    style_header(ws, 1, 4)
    labels = [(2, "基準效率提升（書）", "0%"), (3, "培訓時數（書：2–4 小時）", "0"),
              (4, "前兩週生產力下降（書：10–15%）", "0.0%")]
    for row, label, fmt in labels:
        ws.cell(row=row, column=1, value=label)
        for col, scenario in zip(SCENARIO_COLS, SCENARIOS):
            cell = ws[f"{col}{row}"]
            cell.value = scenario[row - 1]
            cell.number_format = fmt
            cell.fill = INPUT_FILL
    for row, label, template, fmt in CALC_ROWS:
        ws.cell(row=row, column=1, value=label)
        for col in SCENARIO_COLS:
            cell = ws[f"{col}{row}"]
            cell.value = template.format(c=col)
            cell.number_format = fmt
    for row in (17, 18):
        ws.cell(row=row, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    for col in SCENARIO_COLS:
        ws.column_dimensions[col].width = 16


def build_workbook(path: Path) -> None:
    wb = Workbook()
    build_inputs_sheet(wb)
    build_calc_sheet(wb)
    wb.save(path)
    print(f"已寫出試算表：{path.name}")


def find_soffice() -> str:
    """找 LibreOffice 執行檔；PATH 沒有時找 Windows 預設安裝位置。"""
    found = shutil.which("soffice")
    if found:
        return found
    for candidate in (Path("C:/Program Files/LibreOffice/program/soffice.exe"),
                      Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe")):
        if candidate.exists():
            return str(candidate)
    raise FileNotFoundError("找不到 LibreOffice（soffice），無法重算公式")


def recalc_with_libreoffice(path: Path) -> None:
    """用 LibreOffice 開檔重存，讓公式的計算結果寫進檔案。"""
    soffice = find_soffice()
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        # 獨立的使用者設定檔：避免和使用者已開著的 LibreOffice 搶鎖
        profile = (tmp_dir / "profile").as_uri()
        out_dir = tmp_dir / "out"
        cmd = [soffice, f"-env:UserInstallation={profile}", "--headless",
               "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
               "--outdir", str(out_dir), str(path)]
        result = subprocess.run(cmd, capture_output=True, text=True,
                                encoding="utf-8", errors="replace", timeout=180)
        converted = out_dir / path.name
        if result.returncode != 0 or not converted.exists():
            raise RuntimeError(f"LibreOffice 重算失敗：{result.stderr or result.stdout}")
        shutil.copyfile(converted, path)
    print("已用 LibreOffice 重算公式")


def read_results(path: Path) -> dict:
    """讀出試算表的計算結果（data_only 讀快取值）。"""
    wb = load_workbook(path, data_only=True)
    inputs_ws, calc_ws = wb["假設"], wb["計算"]
    inputs = {name: inputs_ws.cell(row=row, column=3).value for row, name, *_ in INPUTS}
    scenarios = {}
    for col, (name, *_rest) in zip(SCENARIO_COLS, SCENARIOS):
        values = {}
        for row in range(2, 19):
            label = calc_ws[f"A{row}"].value
            values[label] = calc_ws[f"{col}{row}"].value
        if values["ROI"] is None:
            raise ValueError(f"情境「{name}」的 ROI 沒有計算值，公式可能沒被重算")
        scenarios[name] = values
    return {"inputs": inputs, "scenarios": scenarios}


def python_reference(inputs: dict) -> dict:
    """不靠試算表，用 Python 再算一次 ROI 與回本月數，供交叉比對。"""
    v = list(inputs.values())
    team, annual, hours, weeks, hi_share, hi_w, mid_share, mid_w, dip_weeks, sub, months = v
    weight = hi_share * hi_w + mid_share * mid_w
    out = {}
    for name, eff, train_h, dip in SCENARIOS:
        monthly_benefit = team * annual / 12 * eff * weight
        monthly_net = monthly_benefit - team * sub
        one_time = team * train_h * annual / hours + team * annual / weeks * dip_weeks * dip
        total_cost = one_time + team * sub * months
        roi = (monthly_benefit * months - total_cost) / total_cost
        out[name] = (roi, one_time / monthly_net)
    return out


def cross_check(results: dict) -> None:
    reference = python_reference(results["inputs"])
    for name, (roi, payback) in reference.items():
        sheet = results["scenarios"][name]
        if abs(sheet["ROI"] - roi) > 1e-9 or abs(sheet["回本月數"] - payback) > 1e-9:
            raise AssertionError(f"情境「{name}」試算表與 Python 計算不一致")
    print("交叉比對通過：試算表公式結果 = Python 獨立計算")


def money(value: float) -> str:
    return f"USD {value:,.0f}"


def write_summary(results: dict, path: Path) -> None:
    """依試算表結果產生給老闆看的一頁摘要。"""
    inp, sc = results["inputs"], results["scenarios"]
    base = sc["基準"]
    rows = "\n".join(
        f"| {name} | {s['基準效率提升（書）']:.0%} | {s['有效效率提升']:.1%} | "
        f"{money(s['一次性導入成本合計'])} | {money(s['每月淨效益'])} | "
        f"{money(s['分析期淨效益'])} | {s['ROI']:.0%} | {s['回本月數']:.2f} 個月 |"
        for name, s in sc.items()
    )
    team, months = inp["團隊開發者人數"], inp["分析期間"]
    text = f"""# Claude Code 導入商業論證（一頁摘要）

> 本頁所有數字由 `build_case.py` 從 `business-case.xlsx` 的公式計算結果讀出，未手抄。
> 改 `假設` 工作表的黃底欄位後重跑腳本即可更新。

## 結論

以 {team} 人團隊、{months} 個月計算，**基準情境 ROI 約 {base['ROI']:.0%}，
約 {base['回本月數']:.1f} 個月回本**；最保守情境 ROI 也有 {sc['保守']['ROI']:.0%}、
{sc['保守']['回本月數']:.1f} 個月回本。建議以小團隊試行一季，用實際數字校正假設。

## 三種情境

| 情境 | 基準效率 | 加權後效率 | 一次性導入成本 | 每月淨效益 | {months} 個月淨效益 | ROI | 回本 |
| --- | --- | --- | --- | --- | --- | --- | --- |
{rows}

## 主要假設

| 項目 | 數值 | 來源 |
| --- | --- | --- |
| 效率提升 保守／基準／樂觀 | 15%／20%／30% | 書 |
| 培訓時數 | 2–4 小時（保守 4、基準 3、樂觀 2） | 書 |
| 前兩週生產力下降 | 10–15%（保守 15%、基準 12.5%、樂觀 10%） | 書 |
| 團隊人數 | {team} 人 | 假設值 |
| 每人年度全額成本 | {money(inp['每人年度全額成本（薪資＋福利＋稅＋設備）'])} | 假設值 |
| 每人每月訂閱 | {money(inp['每人每月訂閱成本'])} | 假設值 |
| 任務加權 | 高影響 {inp['高影響任務佔工時比例']:.0%} × {inp['高影響任務效率權重']}、\
中影響 {inp['中影響任務佔工時比例']:.0%} × {inp['中影響任務效率權重']}（加權係數 {base['加權係數（高／中影響任務）']:.2f}） | 假設值 |

## 風險與注意

- 效益是「省下的工時換算成本」，不等於直接減少支出；要靠交付更多功能或縮短交期兌現。
- 標「假設值」的數字書上沒有，導入前請換成公司實際薪資、訂閱方案與任務組成。
- 生產力下降只算前 {inp['生產力下降期間']} 週；若團隊導入較慢，請拉長期間重算。
"""
    path.write_text(text, encoding="utf-8")
    print(f"已寫出摘要：{path.name}")


def main() -> None:
    try:
        build_workbook(XLSX_PATH)
        recalc_with_libreoffice(XLSX_PATH)
        results = read_results(XLSX_PATH)
        cross_check(results)
        write_summary(results, MD_PATH)
    except (OSError, RuntimeError, ValueError, AssertionError,
            subprocess.TimeoutExpired, KeyError) as exc:
        print(f"錯誤：{exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
